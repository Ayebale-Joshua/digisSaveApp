import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

class ExcelGenerator:
    @staticmethod
    def export_meeting_ledger_to_excel(meeting_data: dict, transactions: list, filepath: str) -> str:
        """
        Exports meeting ledger transactions to a formatted Excel workbook.
        """
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = f"Meeting_{meeting_data.get('meeting_number')}"
        
        # Ensure gridlines are visible
        ws.views.sheetView[0].showGridLines = True

        # Header Styling
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="2B6CB0", end_color="2B6CB0", fill_type="solid")
        center_alignment = Alignment(horizontal="center", vertical="center")
        
        # Title Block
        ws.merge_cells("A1:E1")
        ws["A1"] = f"DIGISAVE - MEETING #{meeting_data.get('meeting_number')} ({meeting_data.get('meeting_date')})"
        ws["A1"].font = Font(name="Calibri", size=14, bold=True, color="1A365D")
        ws["A1"].alignment = center_alignment

        # Table Headers
        headers = ["Passbook Number", "Member Name", "Transaction Type", "Fund Source", "Amount (UGX)"]
        ws.append([]) # Blank row
        ws.append(headers)

        for col_num in range(1, 6):
            cell = ws.cell(row=3, column=col_num)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center_alignment

        # Data Rows
        for tx in transactions:
            ws.append([
                tx.get("passbook_number"),
                tx.get("full_name"),
                tx.get("transaction_type"),
                tx.get("fund_source"),
                tx.get("amount", 0.0)
            ])

        # Save workbook
        wb.save(filepath)
        return filepath